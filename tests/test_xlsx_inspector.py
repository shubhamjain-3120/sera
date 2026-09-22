import hashlib
import io
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

from app.inspectors.xlsx import _list_extended_validations, inspect_xlsx, read_sheet_grid
from app.xlsx_preserve import patch_cell


def workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Application"
    sheet.append(["Name", "Class", "Total"])
    sheet.append(["Existing", "A", "=1+1"])
    sheet.append(["Other", "B", "=2+2"])
    sheet.add_table(Table(displayName="Applicants", ref="A1:C3"))
    validation = DataValidation(type="list", formula1='"A,B,C"')
    sheet.add_data_validation(validation)
    validation.add("B2:B1048576")
    hidden = workbook.create_sheet("Lookups")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "A"
    workbook.defined_names.add(DefinedName("ApplicantName", attr_text="Application!$A$2"))
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_inspector_exposes_structure_without_expanding_whole_columns():
    data = workbook_bytes()
    result = inspect_xlsx(io.BytesIO(data))
    assert len(result["fields"]) == 1
    assert result["fields"][0]["location"]["cell_range"] == "B2:B1048576"
    assert result["inspection"]["sheets"][1]["state"] == "hidden"
    assert result["inspection"]["sheets"][0]["tables"] == ["Applicants"]
    assert result["inspection"]["defined_names"][0]["name"] == "ApplicantName"


def test_grid_is_explicitly_bounded_and_marks_formulas():
    grid = read_sheet_grid(io.BytesIO(workbook_bytes()), "Application", 1, 3, 1, 3)
    assert len(grid["rows"]) == 3
    assert grid["rows"][1][2]["formula"] is True


def test_preservation_spike_changes_only_target_worksheet_part():
    before = workbook_bytes()
    after = patch_cell(before, "Application", "A2", "Changed")
    with zipfile.ZipFile(io.BytesIO(before)) as old, zipfile.ZipFile(io.BytesIO(after)) as new:
        changed = [
            name
            for name in old.namelist()
            if hashlib.sha256(old.read(name)).digest() != hashlib.sha256(new.read(name)).digest()
        ]
    assert changed == ["xl/worksheets/sheet1.xml"]
    reopened = load_workbook(io.BytesIO(after), data_only=False)
    assert reopened["Application"]["A2"].value == "Changed"
    assert reopened["Application"]["C2"].value == "=1+1"
    assert reopened["Lookups"].sheet_state == "hidden"
    assert "Applicants" in reopened["Application"].tables
    assert "ApplicantName" in reopened.defined_names
    assert str(reopened["Application"].data_validations.dataValidation[0].sqref) == "B2:B1048576"


def test_preservation_spike_rejects_formula_replacement():
    try:
        patch_cell(workbook_bytes(), "Application", "C2", "not allowed")
    except ValueError as error:
        assert "Formula replacement" in str(error)
    else:
        raise AssertionError("Expected formula edit to be rejected")


def test_x14_validation_reader_keeps_range_formula_and_broken_reference():
    xml = b'''<worksheet xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main"><extLst><x14:dataValidations><x14:dataValidation type="list"><x14:formula1><xm:f>#REF!</xm:f></x14:formula1><xm:sqref>W1001:X1048576</xm:sqref></x14:dataValidation></x14:dataValidations></extLst></worksheet>'''
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as target:
        target.writestr("xl/worksheets/sheet1.xml", xml)
    with zipfile.ZipFile(io.BytesIO(archive.getvalue())) as source:
        found = _list_extended_validations(source)
    assert found == [{"part": "xl/worksheets/sheet1.xml", "range": "W1001:X1048576", "kind": "extended", "type": "list", "formula1": "#REF!"}]
