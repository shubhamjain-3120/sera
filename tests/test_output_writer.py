from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from pypdf import PdfReader
from reportlab.pdfgen import canvas

from app.output_writer import FormWriteError, write_form


def _flat_pdf() -> bytes:
    stream = BytesIO()
    document = canvas.Canvas(stream, pagesize=(300, 300))
    document.drawString(20, 250, "Applicant:")
    document.save()
    return stream.getvalue()


def test_manual_pdf_field_is_visible_in_output():
    field = {
        "id": "applicant",
        "field_type": "text",
        "location": {"kind": "pdf_rect", "page": 1, "rect": [80, 240, 250, 265]},
        "current_value": None,
    }
    output = write_form(_flat_pdf(), {"fields": [field]}, {"applicant": "Client Name"}, "pdf")
    assert "Client Name" in PdfReader(BytesIO(output)).pages[0].extract_text()


def test_pdf_action_is_not_written():
    field = {
        "id": "reset", "field_type": "action", "writable": False,
        "location": {"kind": "pdf_rect", "page": 1, "rect": [80, 240, 250, 265]},
        "current_value": None,
    }
    original = _flat_pdf()
    assert write_form(original, {"fields": [field]}, {"reset": "Wrong"}, "pdf") == original


def _workbook() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Application"
    sheet["A1"] = "Previous"
    sheet["B2"] = 4
    sheet["B3"].number_format = "mm/dd/yyyy"
    sheet["C3"] = "=B2+1"
    validation = DataValidation(type="list", formula1='"Yes,No"')
    validation.add("A1")
    sheet.add_data_validation(validation)
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def test_workbook_writes_reviewed_cells_and_preserves_formula_and_dropdown():
    fields = [
        {"id": "answer", "field_type": "choice", "location": {"sheet": "Application", "cell_range": "A1"}, "current_value": "Previous"},
        {"id": "count", "field_type": "number", "location": {"sheet": "Application", "cell_range": "B2"}, "current_value": 4},
        {"id": "new", "field_type": "text", "location": {"sheet": "Application", "cell_range": "D4"}, "current_value": None},
        {"id": "birthday", "field_type": "date", "location": {"sheet": "Application", "cell_range": "B3"}, "current_value": None},
    ]
    output = write_form(
        _workbook(), {"fields": fields}, {"answer": "Yes", "count": 12, "new": "Added", "birthday": "01-26-1992"}, "xlsx"
    )
    sheet = load_workbook(BytesIO(output))["Application"]
    assert (sheet["A1"].value, sheet["B2"].value, sheet["C3"].value, sheet["D4"].value) == (
        "Yes", 12, "=B2+1", "Added"
    )
    assert len(sheet.data_validations.dataValidation) == 1
    assert sheet["B3"].value.year == 1992


def test_formula_write_fails_with_field_id():
    field = {"id": "formula", "field_type": "number", "location": {"sheet": "Application", "cell_range": "C3"}, "current_value": None}
    with pytest.raises(FormWriteError) as error:
        write_form(_workbook(), {"fields": [field]}, {"formula": 9}, "xlsx")
    assert error.value.field_id == "formula"
