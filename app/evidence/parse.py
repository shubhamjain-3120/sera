import csv
import io
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import Any

import pdfplumber
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from app.parsers.base import CanonicalParse

NATIVE_PARSER_VERSION = "native-evidence-v1"


def _pdf_blocks(content: bytes) -> list[dict[str, Any]]:
    reader = PdfReader(io.BytesIO(content))
    blocks: list[dict[str, Any]] = []
    with pdfplumber.open(io.BytesIO(content)) as document:
        for page_number, page in enumerate(reader.pages, 1):
            width, height = float(page.mediabox.width), float(page.mediabox.height)
            rotation = int(page.get("/Rotate", 0) or 0)
            words = document.pages[page_number - 1].extract_words(
                x_tolerance=3, y_tolerance=3, keep_blank_chars=False
            )
            page_blocks = _word_blocks(words, page_number, width, height, rotation)
            if page_blocks:
                blocks.extend(page_blocks)
            else:
                page_blocks = _ocr_pdf_page(content, page_number, width, height, rotation)
                blocks.extend(page_blocks or [_unreadable_pdf_page(page_number, width, height, rotation)])
    return blocks


def _word_blocks(
    words: list[dict[str, Any]], page: int, width: float, height: float, rotation: int
) -> list[dict[str, Any]]:
    lines: list[list[dict[str, Any]]] = []
    for word in sorted(words, key=lambda item: (float(item["top"]), float(item["x0"]))):
        if not lines or abs(float(lines[-1][0]["top"]) - float(word["top"])) > 3:
            lines.append([word])
        else:
            lines[-1].append(word)
    blocks = []
    for line in lines:
        ordered = sorted(line, key=lambda item: float(item["x0"]))
        top = min(float(word["top"]) for word in ordered)
        bottom = max(float(word["bottom"]) for word in ordered)
        blocks.append(
            {
                "type": "text",
                "text": " ".join(str(word["text"]) for word in ordered),
                "source": {
                    "kind": "pdf_rect",
                    "page": page,
                    "rect": [
                        min(float(word["x0"]) for word in ordered),
                        height - bottom,
                        max(float(word["x1"]) for word in ordered),
                        height - top,
                    ],
                    "coordinate_system": "pdf-points-bottom-left",
                    "rotation": rotation,
                    "rotation_transform": _rotation_transform(rotation, width, height),
                    "page_width": width,
                    "page_height": height,
                },
            }
        )
    return blocks


def _rotation_transform(rotation: int, width: float, height: float) -> list[float]:
    normalized = rotation % 360
    if normalized == 90:
        return [0, 1, -1, 0, height, 0]
    if normalized == 180:
        return [-1, 0, 0, -1, width, height]
    if normalized == 270:
        return [0, -1, 1, 0, 0, width]
    return [1, 0, 0, 1, 0, 0]


def _xlsx_blocks(content: bytes) -> list[dict[str, Any]]:
    _validate_xlsx_package(content)
    workbook = load_workbook(io.BytesIO(content), data_only=False, read_only=False)
    blocks: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                blocks.append(
                    {
                        "type": "cell",
                        "text": str(cell.value),
                        "source": {
                            "kind": "xlsx_range",
                            "sheet": sheet.title,
                            "cell_range": cell.coordinate,
                            "coordinate_system": "excel-a1",
                        },
                    }
                )
    return blocks


def _validate_xlsx_package(content: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > 20_000:
            raise ValueError("Workbook package has too many entries")
        if sum(entry.file_size for entry in entries) > 200 * 1024 * 1024:
            raise ValueError("Workbook expands beyond the safe inspection limit")
        if any(".." in Path(entry.filename).parts or entry.filename.startswith("/") for entry in entries):
            raise ValueError("Workbook package contains an unsafe part name")


def _image_blocks(content: bytes) -> list[dict[str, Any]]:
    with Image.open(io.BytesIO(content)) as image:
        width, height = image.size
        image.load()
        blocks = _ocr_image(image, "image_rect", 1, float(width), float(height), 1.0, 0)
    return blocks or [
        {
            "type": "unreadable",
            "text": "",
            "source": {
                "kind": "image_rect",
                "page": 1,
                "rect": [0.0, 0.0, float(width), float(height)],
                "coordinate_system": "pixels-top-left",
                "rotation": 0,
                "rotation_transform": [1, 0, 0, 1, 0, 0],
                "page_width": float(width),
                "page_height": float(height),
            },
            "reason": "Image requires OCR provider; content was not guessed",
        }
    ]


def _unreadable_pdf_page(page: int, width: float, height: float, rotation: int) -> dict[str, Any]:
    return {
        "type": "unreadable",
        "text": "",
        "source": {
            "kind": "pdf_rect",
            "page": page,
            "rect": [0.0, 0.0, width, height],
            "coordinate_system": "pdf-points-bottom-left",
            "rotation": rotation,
            "rotation_transform": _rotation_transform(rotation, width, height),
            "page_width": width,
            "page_height": height,
        },
        "reason": "No machine-readable text and local OCR is unavailable or unreadable",
    }


def _ocr_pdf_page(
    content: bytes, page: int, width: float, height: float, rotation: int
) -> list[dict[str, Any]]:
    if shutil.which("tesseract") is None:
        return []
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(content)
    image = document[page - 1].render(scale=2.0, draw_annots=True).to_pil()
    return _ocr_image(image, "pdf_rect", page, width, height, 2.0, rotation)


def _ocr_image(
    image: Image.Image,
    kind: str,
    page: int,
    original_width: float,
    original_height: float,
    scale: float,
    rotation: int,
) -> list[dict[str, Any]]:
    if shutil.which("tesseract") is None:
        return []
    blocks: list[dict[str, Any]] = []
    image_width, image_height = image.size
    passes = [(image, 0.0, 0.0, 1.0, 1.0)]
    detail_box = (
        int(image_width * 0.32),
        int(image_height * 0.22),
        int(image_width * 0.70),
        int(image_height * 0.58),
    )
    detail_width = detail_box[2] - detail_box[0]
    detail_height = detail_box[3] - detail_box[1]
    detail = image.crop(detail_box).resize((1800, 1000))
    passes.append(
        (
            detail,
            float(detail_box[0]),
            float(detail_box[1]),
            1800 / detail_width,
            1000 / detail_height,
        )
    )
    for pass_image, offset_x, offset_y, zoom_x, zoom_y in passes:
        with tempfile.NamedTemporaryFile(suffix=".png") as source_file:
            pass_image.save(source_file.name, format="PNG")
            try:
                result = subprocess.run(
                    ["tesseract", source_file.name, "stdout", "--psm", "6", "tsv"],
                    check=True,
                    capture_output=True,
                    text=True,
                    timeout=60,
                )
            except (subprocess.CalledProcessError, subprocess.TimeoutExpired):
                continue
        grouped: dict[tuple[str, str, str, str], list[dict[str, str]]] = {}
        for word in csv.DictReader(io.StringIO(result.stdout), delimiter="\t"):
            if not (word.get("text") or "").strip():
                continue
            line_key = tuple(word.get(name, "") for name in ("page_num", "block_num", "par_num", "line_num"))
            grouped.setdefault(line_key, []).append(word)
        for words in grouped.values():
            confidences = [float(word["conf"]) for word in words if float(word.get("conf", -1)) >= 0]
            confidence = sum(confidences) / len(confidences) if confidences else 0
            left = min(int(word["left"]) for word in words) / zoom_x + offset_x
            top = min(int(word["top"]) for word in words) / zoom_y + offset_y
            right = max(int(word["left"]) + int(word["width"]) for word in words) / zoom_x + offset_x
            bottom = max(int(word["top"]) + int(word["height"]) for word in words) / zoom_y + offset_y
            if kind == "pdf_rect":
                rect = [left / scale, original_height - bottom / scale, right / scale, original_height - top / scale]
                coordinate_system = "pdf-points-bottom-left"
            else:
                rect = [left, top, right, bottom]
                coordinate_system = "pixels-top-left"
            source = {
                "kind": kind,
                "page": page,
                "rect": rect,
                "coordinate_system": coordinate_system,
                "rotation": rotation,
                "rotation_transform": _rotation_transform(rotation, original_width, original_height),
                "page_width": original_width,
                "page_height": original_height,
            }
            text = " ".join(word["text"].strip() for word in words)
            if confidence < 45:
                blocks.append({"type": "unreadable", "text": text, "source": source, "reason": f"Low OCR confidence ({confidence:.0f}%)"})
            else:
                blocks.append({"type": "text", "text": text, "source": source, "ocr_confidence": confidence / 100})
    return blocks


def _text_blocks(content: bytes) -> list[dict[str, Any]]:
    text = content.decode("utf-8", errors="replace")
    blocks: list[dict[str, Any]] = []
    offset = 0
    for line in text.splitlines(keepends=True):
        value = line.rstrip("\r\n")
        if value.strip():
            blocks.append(
                {
                    "type": "text",
                    "text": value,
                    "source": {
                        "kind": "text_span",
                        "char_start": offset,
                        "char_end": offset + len(value),
                        "coordinate_system": "unicode-codepoints",
                    },
                }
            )
        offset += len(line)
    return blocks


def native_parse(kind: str, content: bytes) -> CanonicalParse:
    if kind == "pdf":
        blocks = _pdf_blocks(content)
    elif kind == "xlsx":
        blocks = _xlsx_blocks(content)
    elif kind == "image":
        blocks = _image_blocks(content)
    elif kind == "text":
        blocks = _text_blocks(content)
    else:
        raise ValueError(f"Unsupported source kind: {kind}")
    return CanonicalParse(
        blocks=blocks,
        raw={"provider": "native", "block_count": len(blocks)},
        provider_job_id=None,
        parser_version=NATIVE_PARSER_VERSION,
    )
